# -*- coding: utf-8 -*-
"""
Exporteer alle Firefly-prompts uit cards.json naar één Word-document (v3-stijl).
Ontbrekende prompts (HEROES / DL1) worden aangevuld met voorstelteksten.
"""
from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
CARDS_PATH = ROOT / "src/data/gesprekskaarten/cards.json"
OUT_REPO = ROOT / "docs/gesprekskaarten/Firefly_prompts_alle_58_website.docx"
OUT_ONEDRIVE = Path(
    r"C:\Users\876409\OneDrive - Office 365 Fontys\Gesprekskaarten"
) / "Firefly_prompts_alle_58_website.docx"

FIREFLY_SUFFIX = (
    "Photorealistic, editorial photography style, natural lighting with slight drama, "
    "shallow depth of field, real people in authentic settings, no text, no graphics, "
    "no illustrations, cinematic composition, suitable for A5 print, portrait orientation 3:4, "
    "bottom 20% slightly dark or uncluttered for text overlay."
)

# Voorstelprompts waar de site (nog) geen prompt in JSON heeft
SUPPLEMENTAL: dict[str, str] = {
    "GK_DL_01": (
        "Inside a crowded commuter train, a young passenger hesitates with a soft scarf in hand, "
        "glancing at a teenage girl lying on the floor shivering, conductor in the background, "
        "tense compassionate moment, cool interior lighting"
    ),
    "GK_NM_01": (
        "A nuclear medicine technologist in scrubs gently positioning an anxious middle-aged woman "
        "on a PET scan table, hospital room with imaging equipment, compassionate but hurried atmosphere, "
        "clinical soft lighting"
    ),
    "GK_NM_02": (
        "An elderly man sitting in a hospital scanning room shaking his head while a male technologist "
        "pulls up a chair to listen, PET/CT equipment visible, quiet emotional tension, hospital lighting"
    ),
    "GK_NM_03": (
        "A pediatric nurse kneeling beside a pale young girl in a hospital bed, blood draw supplies nearby, "
        "gentle intimate moment, muted clinical light, emotional restraint"
    ),
    "GK_NM_04": (
        "A paramedic kneeling beside a gurney during active CPR, attempting a blood draw in a brief pause "
        "between compressions, chaotic emergency room, harsh clinical lighting, extreme focus"
    ),
    "GK_RI_01": (
        "A young female researcher at a university desk facing her senior supervisor across papers, "
        "laptop showing an author list, power imbalance visible in body language, office daylight"
    ),
    "GK_RI_02": (
        "Academic journal editors around a conference table reading an email on a laptop, tense silence, "
        "stack of manuscripts, university office, overcast window light"
    ),
    "GK_RI_03": (
        "A senior researcher alone at a laptop preparing another grant application, conference badges "
        "and past presentation folders on the desk, weary conflicted expression, evening office lamp"
    ),
}

VERMELDING = (
    "Indeling volgens de driedeling micro – meso – macro (werkdefinitie Lectoraat Ethisch Werken, "
    "Fontys; geïnspireerd op het werk van Meijer, 2023)."
)


def sort_key(card: dict) -> tuple:
    m = re.match(r"GK_([A-Z]+)_(\d+)", card["id"])
    if not m:
        return (card["id"], 0)
    return (m.group(1), int(m.group(2)))


def full_prompt(raw: str | None, card_id: str) -> tuple[str, str]:
    """Returns (prompt, bron_note)."""
    if raw and raw.strip():
        text = raw.strip()
        if "portrait orientation 3:4" not in text:
            text = f"{text.rstrip('.')}. {FIREFLY_SUFFIX}"
        return text, "cards.json"
    draft = SUPPLEMENTAL.get(card_id)
    if draft:
        return f"{draft}. {FIREFLY_SUFFIX}", "voorstel (nog niet in cards.json)"
    return "[PROMPT ONTBREEKT — handmatig aanvullen]", "ontbreekt"


def image_filename(card: dict) -> str:
    path = card.get("assets", {}).get("afbeelding") or ""
    return Path(path).name if path else f"{card['id']}.png"


def build_doc(cards: list[dict], out_path: Path) -> None:
    doc = Document()
    doc.add_heading("Firefly-prompts — alle 58 gesprekskaarten (website)", level=0)
    doc.add_paragraph(
        f"Gegenereerd op {date.today().isoformat()} uit moreelvakmanschap.nl-data (cards.json). "
        "Gebruik in Adobe Firefly: Photo, Editorial/Documentary, portrait 3:4. "
        "Sla nieuwe beelden op met de bestandsnaam onder ‘Website-bestand’."
    )
    doc.add_paragraph(
        "Na oplevering: plaats bestanden in public/images/gesprekskaarten/ — "
        "koppeling aan kaarten kan via het kaart-ID (GK_…)."
    )

    uitzonderingen: list[str] = []

    for card in sorted(cards, key=sort_key):
        cid = card["id"]
        titel = card["nl"]["titel"]
        prompt, bron = full_prompt(card.get("assets", {}).get("fireflyPrompt"), cid)
        status = card.get("meta", {}).get("afbeeldingStatus", "oud (tijdelijk)")
        complexiteit = card.get("complexiteit", "meso")

        doc.add_heading(f"{cid} — {titel}", level=1)
        p = doc.add_paragraph()
        p.add_run("Website-bestand: ").bold = True
        p.add_run(image_filename(card))
        doc.add_paragraph(f"Set: {card.get('set', '')} · Complexiteit: {complexiteit}")
        doc.add_paragraph(f"Status afbeelding: {status}")
        doc.add_paragraph(f"Promptbron: {bron}")

        doc.add_paragraph("Firefly-prompt (EN, copy-paste):")
        para = doc.add_paragraph(prompt)
        for run in para.runs:
            run.italic = True
            run.font.size = Pt(10)

        if bron.startswith("voorstel") or bron == "ontbreekt":
            uitzonderingen.append(f"{cid} — {titel}: {bron}")

    doc.add_heading("Overzicht complexiteit", level=1)
    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = "ID"
    hdr[1].text = "Titel"
    hdr[2].text = "Niveau"
    hdr[3].text = "Website-bestand"
    for card in sorted(cards, key=sort_key):
        row = table.add_row().cells
        row[0].text = card["id"]
        row[1].text = card["nl"]["titel"]
        row[2].text = card.get("complexiteit", "")
        row[3].text = image_filename(card)

    doc.add_paragraph(VERMELDING)

    doc.add_heading("Uitzonderingen / voorstelprompts", level=1)
    if uitzonderingen:
        for line in uitzonderingen:
            doc.add_paragraph(line)
    else:
        doc.add_paragraph("Geen.")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


def main() -> None:
    cards = json.loads(CARDS_PATH.read_text(encoding="utf-8"))
    build_doc(cards, OUT_REPO)
    build_doc(cards, OUT_ONEDRIVE)
    print(f"Geschreven:\n  {OUT_REPO}\n  {OUT_ONEDRIVE}")
    print(f"Kaarten: {len(cards)}")


if __name__ == "__main__":
    main()
