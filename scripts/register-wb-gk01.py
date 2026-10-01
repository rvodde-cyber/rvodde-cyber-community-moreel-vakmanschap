"""Append WB-GK01-HBO-NL to 08_werkbladenregister.xlsx (OneDrive)."""
from __future__ import annotations

from pathlib import Path

import openpyxl

XLSX = Path(
    r"C:\Users\876409\OneDrive - Office 365 Fontys\Claude outputs\Educatieve werkvormen\08_werkbladenregister.xlsx"
)

ROW = {
    "Werkblad-ID": "WB-GK01-HBO-NL",
    "Titel werkblad": "Van afwegen naar handelen",
    "Taal": "NL",
    "Niveau": "HBO",
    "Opleiding / beroepscontext": "HRM — gevorderden (vanaf jaar 2)",
    "Complexiteit": "n.v.t. (casusonafhankelijk)",
    "Fase(n) MMV": "Wegen; Handelen",
    "Werkvorm-ID": "MV_61",
    "Oorspronkelijke werkvorm": "eigen werk (Lectoraat Ethisch Werken)",
    "Oorspronkelijke bron (APA)": "n.v.t.",
    "Licentielabel bron": "n.v.t.",
    "Naamsvermelding (verplicht bij B)": "",
    "Licentie werkblad": "CC BY-SA 4.0",
    "Versie": "1.0",
    "Datum": "2026-10-01",
    "Status": "Klaar (sfeerbeeld ingevoegd)",
    "URL moreelvakmanschap.nl": "https://moreelvakmanschap.nl/bibliotheek/werkbladen#MV_61",
    "Opmerkingen": "Registercode WB-GK01-HBO-NL; site-id MV_61",
}


def main() -> None:
    if not XLSX.exists():
        raise SystemExit(f"missing: {XLSX}")
    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == ROW["Werkblad-ID"]:
            print("already_registered")
            return
    new_row = ws.max_row + 1
    for col, header in enumerate(headers, start=1):
        if header in ROW:
            ws.cell(new_row, col, value=ROW[header])
    wb.save(XLSX)
    print(f"registered_row={new_row}")


if __name__ == "__main__":
    main()
